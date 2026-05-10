"""
generate_report.py — Tạo báo cáo kiểm thử tự động dạng Excel
Chạy: python generate_report.py
Output: reports/test_report.xlsx
"""
from pathlib import Path
from datetime import datetime
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Data ──────────────────────────────────────────────────────────────────────
TEST_DATA = [
    # (TC_ID, Module, Method, Scenario, Expected, Actual, DB_Verify, Status, Severity)
    ("UDN_11",   "Login",           "test_user_login_success",          "Dang nhap hop le",                          "Login thanh cong",        "Thanh cong",                   "DB verified",         "PASS", "-"),
    ("UDN_13",   "Login",           "test_user_login_wrong_password",    "Sai mat khau",                              "Hien thi loi dang nhap",  "Dung expected",                "DB unchanged",        "PASS", "-"),
    ("UDN_17+18","Login",           "test_user_account_lock",            "Khoa account sau 6 lan sai password",       "Account bi khoa",         "Account van ACTIVE",           "status='ACTIVE'",     "FAIL", "High"),
    ("DK_7",     "Register",        "test_register_success",             "Dang ky hop le",                            "User duoc tao",           "Thanh cong",                   "User inserted DB",    "PASS", "-"),
    ("DK_8",     "Register",        "test_register_duplicate_username",  "Username trung",                            "Reject duplicate",        "Dung expected",                "DB unchanged",        "PASS", "-"),
    ("DK_10",    "Register",        "test_register_last_name_numbers",   "Khong cho nhap so vao Ho",                  "Validation error",        "Register thanh cong",          "User inserted DB",    "FAIL", "Medium"),
    ("DK_11",    "Register",        "test_register_last_name_special",   "Khong cho ky tu dac biet vao Ho",           "Validation error",        "Register thanh cong",          "User inserted DB",    "FAIL", "Medium"),
    ("DK_13",    "Register",        "test_register_first_name_numbers",  "Khong cho nhap so vao Ten",                 "Validation error",        "Register thanh cong",          "User inserted DB",    "FAIL", "Medium"),
    ("DK_14",    "Register",        "test_register_first_name_special",  "Khong cho ky tu dac biet vao Ten",          "Validation error",        "Register thanh cong",          "User inserted DB",    "FAIL", "Medium"),
    ("DK_16",    "Register",        "test_register_username_special",    "Username khong chua ky tu dac biet",        "Reject username",         "User %$%@ duoc tao",           "User inserted DB",    "FAIL", "High"),
    ("DK_17",    "Register",        "test_register_empty_email",         "Email khong duoc de trong",                 "Reject request",          "Register thanh cong",          "User inserted DB",    "FAIL", "High"),
    ("DK_18",    "Register",        "test_register_duplicate_email",     "Email trung",                               "Reject duplicate",        "Dung expected",                "DB unchanged",        "PASS", "-"),
    ("DK_19",    "Register",        "test_register_invalid_email",       "Email sai format",                          "Validation error",        "Dung expected",                "DB unchanged",        "PASS", "-"),
    ("DK_20",    "Register",        "test_register_empty_phone",         "Phone khong duoc de trong",                 "Reject request",          "Register thanh cong",          "User inserted DB",    "FAIL", "High"),
    ("DK_21",    "Register",        "test_register_duplicate_phone",     "Phone trung",                               "Reject duplicate",        "Dung expected",                "DB unchanged",        "PASS", "-"),
    ("DK_27",    "Register",        "test_register_future_dob",          "Khong cho DOB tuong lai",                   "Validation error",        "DOB 02/02/2030 accepted",      "User inserted DB",    "FAIL", "Medium"),
    ("DK_32+33", "Register",        "test_register_weak_password",       "Reject password yeu",                       "Validation error",        "Password 123456 accepted",     "User inserted DB",    "FAIL", "Critical"),
    ("DMK_7",    "Change Password", "test_change_password_success",      "Doi mat khau hop le",                       "Password updated",        "Thanh cong",                   "Password hash changed","PASS", "-"),
    ("DMK_15+16","Change Password", "test_change_password_weak",         "Khong cho password yeu",                    "Reject password",         "Password doi thanh cong",      "DB hash changed",     "FAIL", "Critical"),
]

# ── Styles ────────────────────────────────────────────────────────────────────
def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def font(bold=False, color="FF000000", size=11):
    return Font(bold=bold, color=color, size=size, name="Calibri")

def align(h="center", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def border():
    s = Side(style="thin", color="FFAAAAAA")
    return Border(left=s, right=s, top=s, bottom=s)

STATUS_FILL  = {"PASS": "FF92D050", "FAIL": "FFFF4444"}
STATUS_FONT  = {"PASS": "FF1E4620", "FAIL": "FFFFFFFF"}
SEVERITY_FILL= {"Critical": "FF7030A0", "High": "FFFF0000",
                 "Medium":  "FFED7D31", "Low": "FFFFC000", "-": "FFD9D9D9"}
SEVERITY_FONT= {"Critical": "FFFFFFFF", "High": "FFFFFFFF",
                 "Medium":  "FFFFFFFF", "Low": "FF000000", "-": "FF666666"}

COL_WIDTHS = [10, 18, 35, 38, 30, 32, 22, 10, 12]
HEADERS = [
    "TC ID", "Module", "Automation Method", "Test Scenario",
    "Expected Result", "Actual Result", "DB Verification", "Status", "Severity"
]

def build_excel(out_path="reports/test_report.xlsx"):
    Path("reports").mkdir(exist_ok=True)
    wb = openpyxl.Workbook()

    # ── Sheet 1: Detail ───────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Ket qua kiem thu"

    # Title row
    ws.merge_cells("A1:I1")
    ws["A1"] = "BAO CAO KIEM THU TU DONG — SELENIUM + PYTEST"
    ws["A1"].fill = fill("FF1F3864")
    ws["A1"].font = font(bold=True, color="FFFFFFFF", size=15)
    ws["A1"].alignment = align()
    ws.row_dimensions[1].height = 38

    # Sub-title
    ws.merge_cells("A2:I2")
    ws["A2"] = f"Ngay chay: {datetime.now().strftime('%d/%m/%Y %H:%M')}  |  URL: http://localhost:3000  |  Tool: Selenium WebDriver + pytest"
    ws["A2"].fill = fill("FF2E75B6")
    ws["A2"].font = font(color="FFFFFFFF", size=10)
    ws["A2"].alignment = align()
    ws.row_dimensions[2].height = 20

    # Summary row
    total  = len(TEST_DATA)
    passed = sum(1 for r in TEST_DATA if r[7] == "PASS")
    failed = sum(1 for r in TEST_DATA if r[7] == "FAIL")
    ws.merge_cells("A3:I3")
    ws["A3"] = f"Tong: {total} test cases  |  PASS: {passed}  |  FAIL: {failed}  |  Ti le pass: {passed/total*100:.1f}%"
    ws["A3"].fill = fill("FF4472C4")
    ws["A3"].font = font(bold=True, color="FFFFFFFF", size=11)
    ws["A3"].alignment = align()
    ws.row_dimensions[3].height = 22

    # Header
    ws.row_dimensions[4].height = 24
    for col, (h, w) in enumerate(zip(HEADERS, COL_WIDTHS), 1):
        c = ws.cell(row=4, column=col, value=h)
        c.fill  = fill("FF2E4057")
        c.font  = font(bold=True, color="FFFFFFFF", size=11)
        c.alignment = align(wrap=True)
        c.border = border()
        ws.column_dimensions[get_column_letter(col)].width = w

    # Data rows
    for row_i, row in enumerate(TEST_DATA, start=5):
        ws.row_dimensions[row_i].height = 32
        bg = "FFF2F2F2" if row_i % 2 == 0 else "FFFFFFFF"
        for col, val in enumerate(row, 1):
            c = ws.cell(row=row_i, column=col, value=val)
            c.border = border()
            c.alignment = align(
                h="left" if col in (3, 4, 5, 6) else "center",
                v="center", wrap=True
            )
            # Status column (col 8)
            if col == 8:
                c.fill = fill(STATUS_FILL.get(val, "FFD9D9D9"))
                c.font = font(bold=True, color=STATUS_FONT.get(val, "FF000000"), size=11)
            # Severity column (col 9)
            elif col == 9:
                c.fill = fill(SEVERITY_FILL.get(val, "FFD9D9D9"))
                c.font = font(bold=True, color=SEVERITY_FONT.get(val, "FF000000"), size=10)
            # TC ID column (col 1) - bold
            elif col == 1:
                c.fill = fill(bg)
                c.font = font(bold=True, size=10)
            else:
                c.fill = fill(bg)
                c.font = font(size=10)

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:I{4 + len(TEST_DATA)}"

    # ── Sheet 2: Bug Summary ──────────────────────────────────────────────────
    ws2 = wb.create_sheet("Danh sach Bug")

    ws2.merge_cells("A1:F1")
    ws2["A1"] = "DANH SACH BUG PHAT HIEN BOI AUTOMATION TEST"
    ws2["A1"].fill = fill("FFC00000")
    ws2["A1"].font = font(bold=True, color="FFFFFFFF", size=14)
    ws2["A1"].alignment = align()
    ws2.row_dimensions[1].height = 32

    bug_headers = ["STT", "TC ID", "Module", "Test Scenario", "Actual Result", "Severity"]
    bug_widths  = [6, 12, 18, 40, 35, 12]
    for col, (h, w) in enumerate(zip(bug_headers, bug_widths), 1):
        c = ws2.cell(row=2, column=col, value=h)
        c.fill = fill("FFD9001B")
        c.font = font(bold=True, color="FFFFFFFF", size=11)
        c.alignment = align(wrap=True)
        c.border = border()
        ws2.column_dimensions[get_column_letter(col)].width = w
    ws2.row_dimensions[2].height = 22

    bugs = [r for r in TEST_DATA if r[7] == "FAIL"]
    for i, r in enumerate(bugs, 1):
        ws2.row_dimensions[i + 2].height = 34
        bg = "FFF2F2F2" if i % 2 == 0 else "FFFFFFFF"
        row_vals = [i, r[0], r[1], r[3], r[5], r[8]]
        for col, val in enumerate(row_vals, 1):
            c = ws2.cell(row=i + 2, column=col, value=val)
            c.border = border()
            c.alignment = align(h="left" if col in (4, 5) else "center",
                                v="center", wrap=True)
            if col == 6:
                c.fill = fill(SEVERITY_FILL.get(str(val), "FFD9D9D9"))
                c.font = font(bold=True, color=SEVERITY_FONT.get(str(val), "FF000000"), size=10)
            else:
                c.fill = fill(bg)
                c.font = font(bold=(col == 2), size=10)

    ws2.freeze_panes = "A3"

    # ── Sheet 3: Summary Stats ────────────────────────────────────────────────
    ws3 = wb.create_sheet("Tong quan")

    ws3.merge_cells("A1:D1")
    ws3["A1"] = "TONG QUAN KET QUA KIEM THU"
    ws3["A1"].fill = fill("FF1F3864")
    ws3["A1"].font = font(bold=True, color="FFFFFFFF", size=14)
    ws3["A1"].alignment = align()
    ws3.row_dimensions[1].height = 32

    # Stats table
    stats = [
        ("Tong so test cases", total,  "FF4472C4", "FFFFFFFF"),
        ("So test PASS",        passed, "FF70AD47", "FFFFFFFF"),
        ("So test FAIL",        failed, "FFFF4444", "FFFFFFFF"),
        ("Ti le PASS",   f"{passed/total*100:.1f}%", "FF70AD47", "FFFFFFFF"),
    ]
    for row_i, (label, val, bg_clr, fg_clr) in enumerate(stats, start=3):
        ws3.row_dimensions[row_i].height = 44
        ws3.column_dimensions["A"].width = 28
        ws3.column_dimensions["B"].width = 18
        lc = ws3.cell(row=row_i, column=1, value=label)
        lc.fill = fill("FFD6DCE4")
        lc.font = font(bold=True, size=13)
        lc.alignment = align(h="left")
        lc.border = border()
        vc = ws3.cell(row=row_i, column=2, value=val)
        vc.fill = fill(bg_clr)
        vc.font = font(bold=True, color=fg_clr, size=18)
        vc.alignment = align()
        vc.border = border()

    # Module breakdown
    ws3.row_dimensions[8].height = 22
    mod_headers = ["Module", "PASS", "FAIL", "Tong"]
    mod_widths   = [22, 10, 10, 10]
    for col, (h, w) in enumerate(zip(mod_headers, mod_widths), 1):
        c = ws3.cell(row=8, column=col, value=h)
        c.fill = fill("FF2E4057")
        c.font = font(bold=True, color="FFFFFFFF")
        c.alignment = align()
        c.border = border()
        ws3.column_dimensions[get_column_letter(col)].width = w

    modules = {}
    for r in TEST_DATA:
        m = r[1]
        if m not in modules:
            modules[m] = {"PASS": 0, "FAIL": 0}
        modules[m][r[7]] += 1

    for row_i, (mod, cnt) in enumerate(modules.items(), start=9):
        ws3.row_dimensions[row_i].height = 22
        bg = "FFF2F2F2" if row_i % 2 == 0 else "FFFFFFFF"
        tot = sum(cnt.values())
        for col, val in enumerate([mod, cnt["PASS"], cnt["FAIL"], tot], 1):
            c = ws3.cell(row=row_i, column=col, value=val)
            c.fill = fill(bg)
            c.font = font(size=11)
            c.alignment = align(h="left" if col == 1 else "center")
            c.border = border()

    wb.save(out_path)
    print(f"[OK] Saved: {out_path}")

if __name__ == "__main__":
    build_excel()
