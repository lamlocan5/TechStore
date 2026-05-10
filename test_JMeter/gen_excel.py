import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, GradientFill
from openpyxl.utils import get_column_letter
from openpyxl.styles.numbers import FORMAT_NUMBER_COMMA_SEPARATED1

wb = openpyxl.Workbook()

# ── Colors ──────────────────────────────────────────────
C_HEADER    = "1F3864"   # dark navy
C_TITLE     = "2E75B6"   # blue
C_LOAD      = "BDD7EE"   # light blue
C_STRESS    = "FFE699"   # yellow
C_SPIKE     = "F4B942"   # orange
C_PASS      = "C6EFCE"   # green
C_FAIL      = "FFC7CE"   # red
C_WARN      = "FFEB9C"   # light yellow
C_ROW_ALT   = "F2F2F2"   # grey
C_WHITE     = "FFFFFF"
C_BORDER    = "9DC3E6"
C_CRITICAL  = "FF0000"

thin = Side(style="thin", color="BFBFBF")
med  = Side(style="medium", color=C_TITLE)
B_THIN  = Border(left=thin, right=thin, top=thin, bottom=thin)
B_MED   = Border(left=med,  right=med,  top=med,  bottom=med)

def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def hdr(text, size=11, bold=True, color="FFFFFF", wrap=True):
    return {"value": text, "font": Font(bold=bold, size=size, color=color, name="Calibri"),
            "fill": fill(C_HEADER), "alignment": Alignment(horizontal="center", vertical="center", wrap_text=wrap),
            "border": B_THIN}

def cell_style(value, bg=C_WHITE, bold=False, color="000000", align="left", wrap=True):
    return {"value": value, "font": Font(bold=bold, size=10, color=color, name="Calibri"),
            "fill": fill(bg), "alignment": Alignment(horizontal=align, vertical="center", wrap_text=wrap),
            "border": B_THIN}

def write_row(ws, row_idx, data_list):
    for col_idx, d in enumerate(data_list, 1):
        c = ws.cell(row=row_idx, column=col_idx)
        c.value = d["value"]
        c.font = d.get("font", Font())
        c.fill = d.get("fill", PatternFill())
        c.alignment = d.get("alignment", Alignment())
        c.border = d.get("border", Border())

def set_col_widths(ws, widths):
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

def row_height(ws, row, h):
    ws.row_dimensions[row].height = h

# ═══════════════════════════════════════════════════════
# SHEET 1: TỔNG QUAN
# ═══════════════════════════════════════════════════════
ws1 = wb.active
ws1.title = "1. Tổng Quan"
ws1.sheet_view.showGridLines = False

# Title
ws1.merge_cells("A1:H1")
c = ws1["A1"]
c.value = "BÁO CÁO KIỂM THỬ HIỆU NĂNG – BOOKSTORE MICROSERVICES"
c.font = Font(bold=True, size=16, color="FFFFFF", name="Calibri")
c.fill = fill("1F3864")
c.alignment = Alignment(horizontal="center", vertical="center")
row_height(ws1, 1, 40)

ws1.merge_cells("A2:H2")
c = ws1["A2"]
c.value = "Apache JMeter | API Gateway: localhost:8888 | MySQL 8 | 8 Kịch Bản Test"
c.font = Font(italic=True, size=11, color="FFFFFF", name="Calibri")
c.fill = fill("2E75B6")
c.alignment = Alignment(horizontal="center", vertical="center")
row_height(ws1, 2, 22)

row_height(ws1, 3, 8)

# Header row
headers = ["ID", "Tên Kịch Bản", "Loại Test", "Threads", "Ramp-up", "Tổng Request", "Throughput", "Mức Độ Rủi Ro"]
write_row(ws1, 4, [hdr(h) for h in headers])
row_height(ws1, 4, 30)

data = [
    ("TG_01","User Login – Auth Load Test","Load Test",100,"60s",100,"50 req/s","🔴 High"),
    ("TG_02","Product Listing – Homepage Load Test","Load Test",200,"60s",1000,"100 req/s","🔴 High"),
    ("TG_03","Advanced Search – Stress Test","Stress Test",200,"120s",2000,"30 req/s","🔴 Critical"),
    ("TG_04","Product Detail – Concurrent Read","Load Test",150,"60s",1500,"80 req/s","🔴 High"),
    ("TG_05","Add to Cart – Race Condition Spike","Spike Test",100,"5s",300,"50 req/s","🔴 High"),
    ("TG_06","Create Order – Transaction Stress","Stress Test",50,"30s",100,"20 req/s","🔴 Critical"),
    ("TG_07","My Orders – Pagination Load Test","Load Test",100,"60s",500,"40 req/s","🟡 Medium"),
    ("TG_08","Voucher Claim – Race Condition Spike","Spike Test",80,"3s",80,"80 req/s","🔴 High"),
]

TYPE_COLOR = {"Load Test": C_LOAD, "Stress Test": C_STRESS, "Spike Test": C_SPIKE}
RISK_COLOR = {"🔴 Critical": "FFC7CE", "🔴 High": "FFE699", "🟡 Medium": "C6EFCE"}

for i, (tid, name, ttype, threads, ramp, total, tput, risk) in enumerate(data):
    row = 5 + i
    bg = C_ROW_ALT if i % 2 else C_WHITE
    tc = TYPE_COLOR.get(ttype, C_WHITE)
    rc = RISK_COLOR.get(risk, C_WHITE)
    vals = [
        cell_style(tid, bg, bold=True, color=C_TITLE, align="center"),
        cell_style(name, bg, bold=False),
        cell_style(ttype, tc, bold=True, align="center"),
        cell_style(threads, bg, align="center"),
        cell_style(ramp, bg, align="center"),
        cell_style(total, bg, align="center"),
        cell_style(tput, bg, align="center"),
        cell_style(risk, rc, bold=True, align="center"),
    ]
    write_row(ws1, row, vals)
    row_height(ws1, row, 22)

set_col_widths(ws1, [8, 38, 14, 10, 10, 14, 14, 16])

# Legend
ws1.merge_cells("A14:H14")
c = ws1["A14"]
c.value = "📌 Chú thích loại test: Load Test = tải ổn định  |  Stress Test = tăng dần tới breaking point  |  Spike Test = đột ngột (mô phỏng flash sale)"
c.font = Font(italic=True, size=9, color="595959")
c.alignment = Alignment(horizontal="left", vertical="center")

# ═══════════════════════════════════════════════════════
# SHEET 2: CẤU HÌNH CHI TIẾT
# ═══════════════════════════════════════════════════════
ws2 = wb.create_sheet("2. Cấu Hình Chi Tiết")
ws2.sheet_view.showGridLines = False

ws2.merge_cells("A1:K1")
c = ws2["A1"]
c.value = "CẤU HÌNH CHI TIẾT 8 KỊCH BẢN JMETER"
c.font = Font(bold=True, size=14, color="FFFFFF", name="Calibri")
c.fill = fill(C_HEADER)
c.alignment = Alignment(horizontal="center", vertical="center")
row_height(ws2, 1, 35)

headers2 = ["ID","Loại Test","Threads","Ramp-up","Loop","Duration Assert","Throughput\n(req/min)","Auth","CSV Files","Cross-Service","Request Body"]
write_row(ws2, 2, [hdr(h) for h in headers2])
row_height(ws2, 2, 36)

data2 = [
    ("TG_01","Load",100,"60s",1,"2000ms","3000","❌","users.csv","❌","JSON: username, password"),
    ("TG_02","Load",200,"60s",5,"2000ms","6000","❌","Random page 1-10","❌","Query: page, limit=12"),
    ("TG_03","Stress",200,"120s",10,"5000ms","1800","❌","search_params.csv","❌","Query: keyword, price, brand, category"),
    ("TG_04","Load",150,"60s",10,"2000ms","4800","❌","product_ids.csv","❌","Path var: {product_id}"),
    ("TG_05","Spike",100,"5s",3,"3000ms","3000","✅","users.csv\nvariant_ids.csv","❌","JSON: variantId, quantity=1"),
    ("TG_06","Stress",50,"30s",2,"8000ms","1200","✅","users.csv\nvariant_ids.csv\naddress_ids.csv","✅ (Product + Identity)","JSON: items, addressId, paymentMethod, voucherCode"),
    ("TG_07","Load",100,"60s",5,"3000ms","2400","✅","users.csv","❌","Query: page=1, limit=12"),
    ("TG_08","Spike",80,"3s",1,"4000ms","4800","✅","users.csv\nvoucher_ids.csv","❌","(trống – chỉ cần token)"),
]

TYPE_BG = {"Load": C_LOAD, "Stress": C_STRESS, "Spike": C_SPIKE}

for i, row_data in enumerate(data2):
    row = 3 + i
    bg = C_ROW_ALT if i % 2 else C_WHITE
    tc = TYPE_BG.get(row_data[1], C_WHITE)
    cells = [
        cell_style(row_data[0], bg, bold=True, color=C_TITLE, align="center"),
        cell_style(row_data[1], tc, bold=True, align="center"),
        cell_style(row_data[2], bg, align="center"),
        cell_style(row_data[3], bg, align="center"),
        cell_style(row_data[4], bg, align="center"),
        cell_style(row_data[5], C_WARN if "8000" in str(row_data[5]) else bg, align="center"),
        cell_style(row_data[6], bg, align="center"),
        cell_style(row_data[7], C_PASS if row_data[7]=="✅" else bg, align="center"),
        cell_style(row_data[8], bg, wrap=True),
        cell_style(row_data[9], C_FAIL if "✅" in str(row_data[9]) else bg, align="center"),
        cell_style(row_data[10], bg, wrap=True),
    ]
    write_row(ws2, row, cells)
    row_height(ws2, row, 36)

set_col_widths(ws2, [8,10,10,10,8,14,12,8,22,20,38])

# ═══════════════════════════════════════════════════════
# SHEET 3: NGƯỠNG SLA
# ═══════════════════════════════════════════════════════
ws3 = wb.create_sheet("3. Ngưỡng SLA")
ws3.sheet_view.showGridLines = False

ws3.merge_cells("A1:I1")
c = ws3["A1"]
c.value = "NGƯỠNG HIỆU NĂNG (SLA) – PASS / FAIL CRITERIA"
c.font = Font(bold=True, size=14, color="FFFFFF")
c.fill = fill(C_HEADER)
c.alignment = Alignment(horizontal="center", vertical="center")
row_height(ws3, 1, 35)

hdrs3 = ["ID","Tên","p90","p95\n(SLA chính)","p99","Error Rate\n(Max)","Throughput\n(Min)","Duration\nAssert","Ghi chú đặc biệt"]
write_row(ws3, 2, [hdr(h) for h in hdrs3])
row_height(ws3, 2, 36)

sla_data = [
    ("TG_01","User Login","< 400ms","< 500ms","< 1000ms","< 1%","≥ 50 req/s","< 2000ms","bcrypt CPU-intensive"),
    ("TG_02","Product List","< 600ms","< 800ms","< 1500ms","< 0.5%","≥ 100 req/s","< 2000ms","Không có cache → real DB load"),
    ("TG_03","Adv. Search","< 1500ms","< 2000ms","< 3000ms","< 2%","≥ 30 req/s","< 5000ms","Tìm breaking point – xem Response Time Graph"),
    ("TG_04","Product Detail","< 400ms","< 600ms","< 1000ms","< 0.5%","≥ 80 req/s","< 2000ms","Check body contains 'variants'"),
    ("TG_05","Add to Cart","< 700ms","< 1000ms","< 2000ms","< 3%","≥ 50 req/s","< 3000ms","Spike 5s – race condition upsert"),
    ("TG_06","Create Order","< 2000ms","< 3000ms","< 5000ms","< 5%","≥ 20 req/s","< 8000ms","⚠️ Reset stock sau test!\n⚠️ Price Manipulation bug!"),
    ("TG_07","My Orders","< 800ms","< 1000ms","—","< 1%","≥ 40 req/s","< 3000ms","Check composite index (user_id, created_at)"),
    ("TG_08","Voucher Claim","—","< 1500ms","—","~87% (OK!)","—","< 4000ms","Đếm: số 200 = đúng 10\nAssert NOT 500 (không assert 200!)"),
]

for i, row_data in enumerate(sla_data):
    row = 3 + i
    bg = C_ROW_ALT if i % 2 else C_WHITE
    note_bg = "FFC7CE" if "⚠️" in str(row_data[8]) else bg
    err_bg = C_WARN if "87%" in str(row_data[5]) else bg
    cells = [
        cell_style(row_data[0], bg, bold=True, color=C_TITLE, align="center"),
        cell_style(row_data[1], bg, bold=True),
        cell_style(row_data[2], bg, align="center"),
        cell_style(row_data[3], C_PASS, bold=True, align="center"),
        cell_style(row_data[4], bg, align="center"),
        cell_style(row_data[5], err_bg, bold=True, align="center"),
        cell_style(row_data[6], bg, align="center"),
        cell_style(row_data[7], C_WARN, align="center"),
        cell_style(row_data[8], note_bg, wrap=True),
    ]
    write_row(ws3, row, cells)
    row_height(ws3, row, 40)

set_col_widths(ws3, [8,18,12,14,12,14,14,12,36])

# ═══════════════════════════════════════════════════════
# SHEET 4: ASSERTIONS
# ═══════════════════════════════════════════════════════
ws4 = wb.create_sheet("4. Assertions")
ws4.sheet_view.showGridLines = False

ws4.merge_cells("A1:G1")
c = ws4["A1"]
c.value = "CẤU HÌNH ASSERTIONS JMETER – TỪNG KỊCH BẢN"
c.font = Font(bold=True, size=14, color="FFFFFF")
c.fill = fill(C_HEADER)
c.alignment = Alignment(horizontal="center", vertical="center")
row_height(ws4, 1, 35)

hdrs4 = ["ID","Assert: Status Code","Assert: Body Contains","Assert: Duration","JSON Extractor\n(lấy token)","Throughput Timer\n(req/phút)","Lưu ý quan trọng"]
write_row(ws4, 2, [hdr(h) for h in hdrs4])
row_height(ws4, 2, 36)

assert_data = [
    ("TG_01","= 200","'token'","< 2000ms","$.result.token → ACCESS_TOKEN","3000","Chọn 'Text Response' (không phải 'Response Message')"),
    ("TG_02","= 200","'result'","< 2000ms","—","6000","Dùng tab Parameters cho GET (không dùng Body Data)"),
    ("TG_03","= 200","—","< 5000ms","—","1800","Thêm Gaussian Timer 500±200ms; cần Response Time Graph"),
    ("TG_04","= 200","'variants'","< 2000ms","—","4800","${product_id} trong URL path (không phải query param)"),
    ("TG_05","= 200","'result'","< 3000ms","$.result.token → ACCESS_TOKEN","3000","Authorization header đặt TRONG request cart, không ở cấp TG"),
    ("TG_06","= 200","'result'","< 8000ms","$.result.token → ACCESS_TOKEN","1200","Authorization TRONG request orders\nReset stock sau test"),
    ("TG_07","= 200","'result'","< 3000ms","$.result.token → ACCESS_TOKEN","2400","Authorization TRONG request my-orders"),
    ("TG_08","NOT 500\n(không assert 200!)","—","< 4000ms","$.result.token → ACCESS_TOKEN","4800","Body để trống\nSau test: SELECT current_usage PHẢI = 10"),
]

for i, row_data in enumerate(assert_data):
    row = 3 + i
    bg = C_ROW_ALT if i % 2 else C_WHITE
    note_bg = C_FAIL if "Reset" in str(row_data[6]) or "PHẢI" in str(row_data[6]) else bg
    status_bg = C_WARN if "NOT" in str(row_data[1]) else C_PASS
    cells = [
        cell_style(row_data[0], bg, bold=True, color=C_TITLE, align="center"),
        cell_style(row_data[1], status_bg, bold=True, align="center"),
        cell_style(row_data[2], bg, align="center"),
        cell_style(row_data[3], C_WARN, align="center"),
        cell_style(row_data[4], bg if row_data[4]=="—" else C_LOAD, align="center", wrap=True),
        cell_style(row_data[5], bg, align="center"),
        cell_style(row_data[6], note_bg, wrap=True),
    ]
    write_row(ws4, row, cells)
    row_height(ws4, row, 40)

set_col_widths(ws4, [8,16,14,12,26,14,40])

# ═══════════════════════════════════════════════════════
# SHEET 5: PASS/FAIL CRITERIA
# ═══════════════════════════════════════════════════════
ws5 = wb.create_sheet("5. Pass-Fail Criteria")
ws5.sheet_view.showGridLines = False

ws5.merge_cells("A1:E1")
c = ws5["A1"]
c.value = "TIÊU CHÍ ĐÁNH GIÁ KẾT QUẢ – PASS / FAIL"
c.font = Font(bold=True, size=14, color="FFFFFF")
c.fill = fill(C_HEADER)
c.alignment = Alignment(horizontal="center", vertical="center")
row_height(ws5, 1, 35)

hdrs5 = ["Tiêu chí","Ngưỡng PASS","Ngưỡng FAIL","Mức độ","Hành động khi FAIL"]
write_row(ws5, 2, [hdr(h) for h in hdrs5])
row_height(ws5, 2, 28)

pf_data = [
    ("Login p95","≤ 500ms","> 1000ms","High","Kiểm tra bcrypt rounds, DB index users"),
    ("Product List p95","≤ 800ms","> 2000ms","High","Thêm Redis cache, kiểm tra DB index"),
    ("Advanced Search p95","≤ 2000ms","> 5000ms","Critical","Thêm index, optimize JOIN query, xem EXPLAIN"),
    ("Product Detail p95","≤ 600ms","> 1500ms","High","Kiểm tra index FK variant.product_id"),
    ("Add Cart p95","≤ 1000ms","> 3000ms","High","Kiểm tra upsert logic, DB lock"),
    ("Create Order p95","≤ 3000ms","> 8000ms","Critical","Optimize cross-service calls, check deadlock"),
    ("My Orders p95","≤ 1000ms","> 2500ms","Medium","Thêm composite index (user_id, created_at)"),
    ("Voucher Claim: số 200","= đúng 10","≠ 10 (over-claim)","Critical","Thêm distributed lock / pessimistic lock"),
    ("Error rate tổng thể","≤ 1%","> 5%","High","Kiểm tra log, DB connections, service health"),
    ("OOM errors","= 0","Bất kỳ","Critical","Tăng JVM heap, kiểm tra memory leak"),
    ("DB connection pool","Không cạn","Pool exhausted","Critical","Tăng max-active connections"),
    ("Voucher over-claim","= 0","current_usage > 10","Critical","Fix race condition – thêm DB lock"),
    ("Price Manipulation","Không có","Client gửi price tùy ý","Critical","Server PHẢI fetch giá từ Product Service"),
]

MUC_DO_BG = {"Critical": "FFC7CE", "High": "FFE699", "Medium": "C6EFCE"}

for i, row_data in enumerate(pf_data):
    row = 3 + i
    bg = C_ROW_ALT if i % 2 else C_WHITE
    md_bg = MUC_DO_BG.get(row_data[3], bg)
    cells = [
        cell_style(row_data[0], bg, bold=True),
        cell_style(row_data[1], C_PASS, bold=True, align="center"),
        cell_style(row_data[2], C_FAIL, bold=True, align="center"),
        cell_style(row_data[3], md_bg, bold=True, align="center"),
        cell_style(row_data[4], bg, wrap=True),
    ]
    write_row(ws5, row, cells)
    row_height(ws5, row, 28)

set_col_widths(ws5, [28,16,22,12,42])

# ═══════════════════════════════════════════════════════
# SHEET 6: MONITORING
# ═══════════════════════════════════════════════════════
ws6 = wb.create_sheet("6. Monitoring")
ws6.sheet_view.showGridLines = False

ws6.merge_cells("A1:E1")
c = ws6["A1"]
c.value = "MONITORING STRATEGY – THEO DÕI TRONG QUÁ TRÌNH TEST"
c.font = Font(bold=True, size=14, color="FFFFFF")
c.fill = fill(C_HEADER)
c.alignment = Alignment(horizontal="center", vertical="center")
row_height(ws6, 1, 35)

hdrs6 = ["Metric","Tool","Ngưỡng WARNING","Ngưỡng CRITICAL","Hành động"]
write_row(ws6, 2, [hdr(h) for h in hdrs6])
row_height(ws6, 2, 28)

mon_data = [
    ("CPU Server","PerfMon / htop","> 70%","> 90%","Giảm threads, kiểm tra background processes"),
    ("JVM Heap","JVM metrics / jstat","> 80%","> 95%","Tăng -Xmx, kiểm tra memory leak (jstack)"),
    ("DB Connections","MySQL SHOW PROCESSLIST","> 80% pool","Pool exhausted","Tăng max-active, kiểm tra connection leak"),
    ("Response Time p95","JMeter Aggregate Report","> SLA threshold","> 2× SLA","Kiểm tra slow query log, service logs"),
    ("Error Rate","JMeter Summary Report","> 1%","> 5%","Kiểm tra stack trace, DB health"),
    ("TPS (Throughput)","JMeter TPS Graph","< 50% target","< 20% target","Kiểm tra bottleneck, giảm tải"),
    ("DB Slow Queries","MySQL slow_query_log","> 1s","> 5s","Chạy EXPLAIN, thêm index"),
    ("DB Deadlocks","MySQL error.log","Bất kỳ","Liên tục","Fix transaction isolation, sắp xếp lại lock order"),
    ("Response Time Graph","JMeter (Stress Test)","Tăng dần","Tăng vọt (spike)","Xác định breaking point tại số thread nào"),
]

for i, row_data in enumerate(mon_data):
    row = 3 + i
    bg = C_ROW_ALT if i % 2 else C_WHITE
    cells = [
        cell_style(row_data[0], bg, bold=True),
        cell_style(row_data[1], bg),
        cell_style(row_data[2], C_WARN, align="center"),
        cell_style(row_data[3], C_FAIL, align="center"),
        cell_style(row_data[4], bg, wrap=True),
    ]
    write_row(ws6, row, cells)
    row_height(ws6, row, 30)

set_col_widths(ws6, [24,24,18,18,42])

# Save
out = r"C:\Users\Admin\Desktop\PTIT\Y4_T2\QA\SQA\test_JMeter\JMeter_Performance_Report.xlsx"
wb.save(out)
print(f"Done! Saved to: {out}")
